import logging
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path
from typing import List, Optional

logger = logging.getLogger(__name__)


@dataclass
class ReportRow:
    """One job's production figures. Areas are in m²; waste is the unfilled
    share of the produced sheets — the number a print shop actually pays for."""

    date: Optional[datetime]
    name: str
    status: str
    files: int
    sheets: int
    poses: int
    area_m2: float
    avg_fill: float  # % weighted by sheet area
    waste_m2: float


@dataclass
class ProductionReport:
    start: Optional[date]
    end: Optional[date]
    rows: List[ReportRow] = field(default_factory=list)

    @property
    def totals(self) -> dict:
        area = sum(r.area_m2 for r in self.rows)
        waste = sum(r.waste_m2 for r in self.rows)
        return {
            "jobs": len(self.rows),
            "files": sum(r.files for r in self.rows),
            "sheets": sum(r.sheets for r in self.rows),
            "poses": sum(r.poses for r in self.rows),
            "area_m2": area,
            "waste_m2": waste,
            "avg_fill": (1 - waste / area) * 100.0 if area > 0 else 0.0,
        }


def build_production_report(jobs, start: Optional[date] = None,
                            end: Optional[date] = None) -> ProductionReport:
    """Aggregates JobModel rows (sheets eagerly loaded) into per-job
    production figures, optionally restricted to [start, end] (inclusive,
    on the job's creation date)."""
    report = ProductionReport(start=start, end=end)
    for job in jobs:
        job_date = job.created_at.date() if job.created_at else None
        if start is not None and (job_date is None or job_date < start):
            continue
        if end is not None and (job_date is None or job_date > end):
            continue

        area = 0.0
        filled = 0.0
        poses = 0
        for sheet in job.sheets:
            sheet_area = (sheet.width_mm or 0.0) * (sheet.height_mm or 0.0) / 1e6
            area += sheet_area
            filled += sheet_area * (sheet.fill_rate or 0.0) / 100.0
            poses += len(sheet.items or [])

        report.rows.append(
            ReportRow(
                date=job.created_at,
                name=job.name,
                status=job.status,
                files=len(job.source_paths or []),
                sheets=len(job.sheets),
                poses=poses,
                area_m2=area,
                avg_fill=(filled / area * 100.0) if area > 0 else 0.0,
                waste_m2=area - filled,
            )
        )
    report.rows.sort(key=lambda r: (r.date or datetime.min), reverse=True)
    return report


def export_production_xlsx(report: ProductionReport, out_path: Path) -> Path:
    """Writes the report as a formatted Excel sheet (headers, rows, totals)."""
    import openpyxl
    from openpyxl.styles import Font

    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Production"

    period = ""
    if report.start or report.end:
        period = (
            f" du {report.start:%d/%m/%Y}" if report.start else ""
        ) + (f" au {report.end:%d/%m/%Y}" if report.end else "")
    ws.append([f"Rapport de production JELOTIA{period}"])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([])

    headers = ["Date", "Job", "Statut", "Fichiers", "Planches", "Poses",
               "Surface (m²)", "Remplissage (%)", "Chute (m²)"]
    ws.append(headers)
    for cell in ws[ws.max_row]:
        cell.font = Font(bold=True)

    for row in report.rows:
        ws.append([
            row.date.strftime("%d/%m/%Y %H:%M") if row.date else "—",
            row.name, row.status, row.files, row.sheets, row.poses,
            round(row.area_m2, 3), round(row.avg_fill, 1), round(row.waste_m2, 3),
        ])

    totals = report.totals
    ws.append([])
    total_row = ["TOTAL", f"{totals['jobs']} job(s)", "", totals["files"],
                 totals["sheets"], totals["poses"], round(totals["area_m2"], 3),
                 round(totals["avg_fill"], 1), round(totals["waste_m2"], 3)]
    ws.append(total_row)
    for cell in ws[ws.max_row]:
        cell.font = Font(bold=True)

    widths = [17, 34, 10, 9, 9, 8, 13, 15, 11]
    for index, width in enumerate(widths, start=1):
        ws.column_dimensions[chr(64 + index)].width = width

    wb.save(str(out_path))
    return out_path
